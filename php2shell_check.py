#!/usr/bin/env python3
"""
Safe PHP2Shell exposure checker.

Reads a list of website URLs from a text file and checks for PHP version
disclosure. It does NOT exploit CVE-2026-109618 / CVE-2026-109619.

Use only against systems you are authorized to test.
"""

import argparse
import csv
import hashlib
import html
import ipaddress
import re
import socket
import sys
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from functools import lru_cache
from urllib.parse import urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from packaging.version import InvalidVersion, Version

# Patched PHP branches for PHP2Shell:
# CVE-2026-109618 and CVE-2026-109619.
PATCHED_VERSIONS = {
    (8, 2): Version("8.2.35"),
    (8, 3): Version("8.3.36"),
    (8, 4): Version("8.4.27"),
    (8, 5): Version("8.5.12"),
}

# Common safe indicators of PHP version. These are passive checks only.
HEADER_KEYS = ("x-powered-by", "server")
PHP_VERSION_RE = re.compile(
    r"(?:PHP[/ ]|PHP\s*)(\d+\.\d+(?:\.\d+)?(?:[.-][0-9A-Za-z]+)*)",
    re.IGNORECASE,
)
META_GENERATOR_RE = re.compile(
    r'<meta[^>]+name=["\']generator["\'][^>]+content=["\']([^"\']+)',
    re.IGNORECASE,
)
PHPSESSID_RE = re.compile(r"PHPSESSID=", re.IGNORECASE)
PHPINFO_RE = re.compile(
    r"<title>phpinfo\(\)</title>|PHP Version\s*</td>|phpinfo\(\)",
    re.IGNORECASE,
)
TITLE_RE = re.compile(r"<title[^>]*>(.*?)</title>", re.IGNORECASE | re.DOTALL)

# ---------------------------------------------------------------------------
# Passive scoring model for PHP presence confidence.
# Thresholds are documented heuristics, not scientific — see README.
# ---------------------------------------------------------------------------
SCORE_PHPINFO = 100        # phpinfo() output detected in body
SCORE_X_POWERED_BY = 40    # X-Powered-By header contains PHP
SCORE_PHPSESSID = 40       # Set-Cookie: PHPSESSID=...
SCORE_INDEX_PHP_200 = 10   # --probe-path returned HTTP 200
SCORE_DIFFERS_ROOT = 10    # probe response differs from /
SCORE_DIFFERS_RANDOM = 10  # probe response differs from random .php
SCORE_SAME_AS_ROOT = -20   # probe response identical to /
SCORE_SAME_AS_RANDOM = -30 # probe response identical to random .php

CONFIDENCE_VERIFIED = 80   # score >= this  → VERIFIED_PHP
CONFIDENCE_POSSIBLE = 30   # score >= this  → POSSIBLE_PHP
# below            → LIKELY_FALSE_POSITIVE


@dataclass
class Result:
    url: str
    ip: str | None
    status: str
    http_status: int | None
    php_version: str | None
    source: str | None
    assessment: str
    detail: str
    index_php_status: int | None = None
    php_confidence: str = "N/A"
    indicators: list[str] = field(default_factory=list)
    content_hash: str | None = None
    content_length: int | None = None
    content_type: str | None = None
    server_header: str | None = None
    page_title: str | None = None


def normalize_url(raw: str) -> str | None:
    raw = raw.strip()
    if not raw or raw.startswith("#"):
        return None
    if not urlparse(raw).scheme:
        raw = "https://" + raw
    parsed = urlparse(raw)
    if not parsed.netloc:
        return None
    return raw


def expand_cidr(raw: str) -> list[str] | None:
    """If the line is a CIDR range (e.g. 192.0.2.0/24), return one
    https:// and one http:// URL for every usable host address.
    Returns None when the line is not a CIDR range.
    """
    if "/" not in raw:
        return None
    try:
        network = ipaddress.ip_network(raw, strict=False)
    except ValueError:
        return None
    if network.version != 4:
        return None
    hosts = list(network.hosts()) or [network.network_address]
    return [f"{scheme}://{ip}" for ip in hosts for scheme in ("https", "http")]


def expand_urls(raw: str) -> list[str]:
    """Return the URLs to check for one input line.

    CIDR lines (e.g. 192.0.2.0/24) are expanded to https:// and http://
    URLs for every usable host. Lines with an explicit scheme are used
    as-is; bare hostnames or host/path values are expanded to both
    https:// and http://.
    """
    raw = raw.strip()
    if not raw or raw.startswith("#"):
        return []
    cidr_urls = expand_cidr(raw)
    if cidr_urls is not None:
        return cidr_urls
    if urlparse(raw).scheme:
        url = normalize_url(raw)
        return [url] if url else []
    if not urlparse("https://" + raw).netloc:
        return []
    return [f"https://{raw}", f"http://{raw}"]


def extract_php_version(text: str) -> tuple[str | None, str | None]:
    """Return (version, source) from a header value or HTML body."""
    match = PHP_VERSION_RE.search(text)
    if match:
        return match.group(1), "response-header/body"

    generator = META_GENERATOR_RE.search(text)
    if generator:
        gen_match = PHP_VERSION_RE.search(generator.group(1))
        if gen_match:
            return gen_match.group(1), "meta-generator"

    return None, None


def body_hash(text: str) -> str:
    """SHA-256 hex digest of the response body (first 100 KB)."""
    return hashlib.sha256(text[:100_000].encode("utf-8", errors="replace")).hexdigest()


def extract_title(text: str) -> str | None:
    match = TITLE_RE.search(text[:50_000])
    if match:
        title = html.unescape(match.group(1)).strip()
        return title[:200] if title else None
    return None


def assess_version(version: str | None) -> tuple[str, str]:
    if not version:
        return "UNKNOWN", "No PHP version disclosure found; manual inventory required."

    try:
        parsed = Version(version)
    except InvalidVersion:
        return "UNKNOWN", f'Could not parse disclosed PHP version "{version}".'

    branch = (parsed.major, parsed.minor)

    if branch not in PATCHED_VERSIONS:
        return (
            "LIKELY_VULNERABLE_OR_UNSUPPORTED",
            f'PHP {version} is not in a currently patched branch '
            f"(8.2, 8.3, 8.4, or 8.5). Verify the actual server runtime and upgrade path.",
        )

    fixed = PATCHED_VERSIONS[branch]
    if parsed >= fixed:
        return "PATCHED", f"PHP {version} meets or exceeds the fixed release {fixed}."
    return "LIKELY_VULNERABLE", f"PHP {version} is below fixed release {fixed}."


@lru_cache(maxsize=None)
def resolve_host(hostname: str) -> str | None:
    """Resolve a hostname to its first IP address, or None on failure."""
    try:
        infos = socket.getaddrinfo(hostname, None)
    except socket.gaierror:
        return None
    return infos[0][4][0]


def confidence_from_score(score: int) -> str:
    if score >= CONFIDENCE_VERIFIED:
        return "VERIFIED_PHP"
    if score >= CONFIDENCE_POSSIBLE:
        return "POSSIBLE_PHP"
    return "LIKELY_FALSE_POSITIVE"


def check_site(
    session: requests.Session,
    url: str,
    timeout: float,
    probe_path: str | None = None,
    verify: bool = False,
) -> Result:
    ip = resolve_host(urlparse(url).hostname or "")
    try:
        response = session.get(url, timeout=timeout, allow_redirects=True)
    except requests.RequestException as exc:
        return Result(url, ip, "ERROR", None, None, None, "UNREACHABLE", str(exc))

    version = None
    source = None
    indicators: list[str] = []
    score = 0

    body_text = response.text[:100_000]

    # --- Passive header inspection (version disclosure) ---
    for header in HEADER_KEYS:
        value = response.headers.get(header, "")
        found, found_source = extract_php_version(value)
        if found:
            version, source = found, f"{header} header"
            break

    # --- Passive body inspection (version disclosure) ---
    if not version:
        version, source = extract_php_version(body_text)

    # --- PHP presence indicators (passive, no extra requests) ---
    x_powered_by = response.headers.get("x-powered-by", "")
    if x_powered_by and "php" in x_powered_by.lower():
        score += SCORE_X_POWERED_BY
        indicators.append("X-Powered-By contains PHP")

    set_cookie = response.headers.get("set-cookie", "")
    if PHPSESSID_RE.search(set_cookie):
        score += SCORE_PHPSESSID
        indicators.append("PHPSESSID cookie observed")

    if PHPINFO_RE.search(body_text):
        score += SCORE_PHPINFO
        indicators.append("phpinfo() output detected")

    # --- Response fingerprint ---
    r_hash = body_hash(body_text)
    r_length = len(response.content)
    r_type = response.headers.get("content-type", "")
    r_server = response.headers.get("server", "")
    r_title = extract_title(body_text)

    assessment, detail = assess_version(version)

    # --- Optional probe path ---
    index_php_status = None
    probe_hash = None
    if probe_path:
        parsed = urlparse(url)
        probe_url = f"{parsed.scheme}://{parsed.netloc}{probe_path}"
        try:
            probe_resp = session.get(probe_url, timeout=timeout, allow_redirects=True)
            index_php_status = probe_resp.status_code
            probe_hash = body_hash(probe_resp.text[:100_000])
        except requests.RequestException:
            index_php_status = None

        if index_php_status == 200:
            score += SCORE_INDEX_PHP_200
            indicators.append(f"{probe_path} returned HTTP 200")
            assessment = "INDEX_PHP_200"
        else:
            shown = index_php_status if index_php_status is not None else "no response"
            indicators.append(f"{probe_path} returned HTTP {shown}")

        detail = f"{probe_path}: HTTP {index_php_status or 'N/A'}. " + detail

    # --- --verify mode: root + random-path comparison ---
    if verify and probe_path and index_php_status == 200:
        # Compare probe response to the root page (already fetched).
        if probe_hash and probe_hash == r_hash:
            score += SCORE_SAME_AS_ROOT
            indicators.append("probe response identical to root page (catch-all suspected)")
        else:
            score += SCORE_DIFFERS_ROOT
            indicators.append("probe response differs from root page")

        # Random .php path probe to detect catch-all error pages.
        parsed = urlparse(url)
        random_path = f"/__php2shell_{hashlib.md5(url.encode()).hexdigest()[:8]}.php"
        random_url = f"{parsed.scheme}://{parsed.netloc}{random_path}"
        try:
            random_resp = session.get(random_url, timeout=timeout, allow_redirects=True)
            random_hash = body_hash(random_resp.text[:100_000])
        except requests.RequestException:
            random_hash = None

        if random_hash is not None:
            if probe_hash and probe_hash == random_hash:
                score += SCORE_SAME_AS_RANDOM
                indicators.append("probe response identical to random .php (catch-all suspected)")
            else:
                score += SCORE_DIFFERS_RANDOM
                indicators.append("probe response differs from random .php")

    # --- Confidence classification ---
    if indicators:
        php_confidence = confidence_from_score(score)
    else:
        php_confidence = "N/A"

    return Result(
        url=url,
        ip=ip,
        status="OK",
        http_status=response.status_code,
        php_version=version,
        source=source,
        assessment=assessment,
        detail=detail,
        index_php_status=index_php_status,
        php_confidence=php_confidence,
        indicators=indicators,
        content_hash=r_hash,
        content_length=r_length,
        content_type=r_type or None,
        server_header=r_server or None,
        page_title=r_title,
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Passively check websites for disclosed PHP versions and PHP2Shell patch status."
    )
    parser.add_argument(
        "input_file",
        help="Text file with one URL, hostname, or CIDR range (e.g. 192.0.2.0/24) per line",
    )
    parser.add_argument(
        "-o",
        "--output",
        default="php2shell_results.xlsx",
        help="Output file (.xlsx for Excel, .csv for CSV)",
    )
    parser.add_argument("--timeout", type=float, default=10.0, help="HTTP timeout in seconds")
    parser.add_argument(
        "--user-agent",
        default="PHP2Shell-Exposure-Checker/1.0 (authorized security assessment)",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=8,
        help="Number of concurrent checks (keep modest for CIDR scans)",
    )
    parser.add_argument(
        "--probe-path",
        default=None,
        help="Extra path to probe on each target origin (e.g. /index.php); "
        "HTTP 200 is flagged as INDEX_PHP_200",
    )
    parser.add_argument(
        "--verify",
        action="store_true",
        help="Deeper false-positive checks: compare probe response to root page "
        "and a random .php path. Requires --probe-path. Adds 1 extra request per host.",
    )
    args = parser.parse_args()

    probe_path = None
    if args.probe_path:
        probe_path = args.probe_path if args.probe_path.startswith("/") else "/" + args.probe_path

    if args.verify and not probe_path:
        print("--verify requires --probe-path to be set.", file=sys.stderr)
        return 2

    try:
        with open(args.input_file, encoding="utf-8") as handle:
            raw_urls = handle.readlines()
    except OSError as exc:
        print(f"Cannot read input file: {exc}", file=sys.stderr)
        return 2

    urls: list[str] = []
    for line in raw_urls:
        urls.extend(expand_urls(line))

    if not urls:
        print("No valid URLs found in input file.", file=sys.stderr)
        return 2

    session = requests.Session()
    session.headers.update({"User-Agent": args.user_agent})

    results: list[Result] = []
    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        for result in pool.map(
            lambda u: check_site(session, u, args.timeout, probe_path, args.verify), urls
        ):
            results.append(result)
            extra = (
                f" index.php={result.index_php_status}"
                if result.index_php_status is not None
                else ""
            )
            conf = f" [{result.php_confidence}]" if result.php_confidence != "N/A" else ""
            print(
                f"{result.assessment:<32} {result.url} "
                f"({result.ip or 'no DNS'}) "
                f"PHP={result.php_version or 'not disclosed'}{extra}{conf}"
            )

    # Deduplicate: if https:// succeeded for a host, drop its http:// result.
    https_ok_hosts: set[str] = set()
    for r in results:
        if r.url.startswith("https://") and r.status == "OK":
            https_ok_hosts.add(urlparse(r.url).hostname or "")
    results = [
        r for r in results
        if not (
            r.url.startswith("http://")
            and (urlparse(r.url).hostname or "") in https_ok_hosts
        )
    ]

    if args.output.lower().endswith(".csv"):
        write_csv(args.output, results)
    else:
        write_xlsx(args.output, results)

    print(f"\nResults written to {args.output}")
    return 0


CSV_HEADER = [
    "url",
    "ip",
    "status",
    "http_status",
    "php_version",
    "version_source",
    "index_php_status",
    "assessment",
    "php_confidence",
    "indicators",
    "content_hash",
    "content_length",
    "content_type",
    "server_header",
    "page_title",
    "detail",
]

PATCHED_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
VULNERABLE_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FALSE_POSITIVE_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
ASSESSMENT_COLUMN = CSV_HEADER.index("assessment") + 1
CONFIDENCE_COLUMN = CSV_HEADER.index("php_confidence") + 1


def result_row(result: Result) -> list:
    return [
        result.url,
        result.ip,
        result.status,
        result.http_status,
        result.php_version,
        result.source,
        result.index_php_status,
        result.assessment,
        result.php_confidence,
        "; ".join(result.indicators) if result.indicators else "",
        result.content_hash,
        result.content_length,
        result.content_type,
        result.server_header,
        result.page_title,
        result.detail,
    ]


def write_csv(path: str, results: list[Result]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(CSV_HEADER)
        for result in results:
            writer.writerow(result_row(result))


def write_xlsx(path: str, results: list[Result]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "results"
    sheet.append(CSV_HEADER)
    sheet.freeze_panes = "A2"

    for result in results:
        sheet.append(result_row(result))
        row = sheet.max_row

        # Color the assessment column.
        cell = sheet.cell(row=row, column=ASSESSMENT_COLUMN)
        if cell.value == "PATCHED":
            cell.fill = PATCHED_FILL
        elif cell.value and "VULNERABLE" in cell.value:
            cell.fill = VULNERABLE_FILL

        # Color the confidence column.
        conf_cell = sheet.cell(row=row, column=CONFIDENCE_COLUMN)
        if conf_cell.value == "VERIFIED_PHP":
            conf_cell.fill = PATCHED_FILL
        elif conf_cell.value == "LIKELY_FALSE_POSITIVE":
            conf_cell.fill = FALSE_POSITIVE_FILL
        elif conf_cell.value == "POSSIBLE_PHP":
            conf_cell.fill = VULNERABLE_FILL

    for index, header in enumerate(CSV_HEADER, start=1):
        width = max(
            [len(header)]
            + [len(str(row[index - 1] or "")) for row in (result_row(r) for r in results)]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 80)

    workbook.save(path)


if __name__ == "__main__":
    raise SystemExit(main())
